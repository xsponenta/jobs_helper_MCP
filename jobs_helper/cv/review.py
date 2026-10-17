"""CV review that mirrors how a CV is actually screened.

1. ATS parse      what an applicant-tracking system extracts from the file (contacts, positions, dates, skills)
2. Knock-outs     hard filters of a vacancy: required years (computed from the dates), title, keywords, English
3. 6-second scan  what a recruiter sees in the top third: target title, years, key skills, a strong number
4. Content        quantified achievements, strong verbs, no duties / clichés / passive voice, timeline gaps
5. Design         measured on the PDF: font sizes, number of fonts, margins, text density, page fill

Every issue carries the number of overall points its fix is worth; the levels Bronze → Silver → Gold → Platinum
show how far the CV is from standing out. Rules only — no LLM — so the score is repeatable and explainable; the
model judges the visual look from the page image on top of this.
"""

from __future__ import annotations

import os
import re
from datetime import date
from pathlib import Path

from jobs_helper.cv.skills import CLICHES, ROLE_WORDS, SECTION_ALIASES, TECH, WEAK_STARTS

PENALTY = {"critical": 40, "major": 18, "minor": 7}
WEIGHTS = {"ats": 0.18, "completeness": 0.12, "first_impression": 0.15, "impact": 0.22, "readability": 0.08,
           "design": 0.10, "match": 0.15}
LEVELS = [(90, "Platinum", "stands out: ready for top companies"), (80, "Gold", "strong; polish the details"),
          (65, "Silver", "solid base, several fixes will make it noticeable"), (0, "Bronze", "needs work before applying")]

_YEAR = re.compile(r"\b(19|20)\d{2}\b")
_NUMBER = re.compile(r"\d")
_METRIC = re.compile(r"\d+(?:[.,]\d+)?\s*(%|x\b|×|k\b|m\b|тис|млн|\$|usd|users|клієнт|customers|hours|год|days|днів|people|людей|engineers)", re.I)
_BULLET = re.compile(r"^\s*([-*•▪●◦]|\d+[.)])\s+")
_MONTHS = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6, "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
    "січ": 1, "лют": 2, "бер": 3, "кві": 4, "тра": 5, "чер": 6, "лип": 7, "сер": 8, "вер": 9, "жов": 10, "лис": 11, "гру": 12,
}
_MON = r"(?:[A-Za-zА-Яа-яІіЇїЄєҐґ]{3,10}\.?\s+)"
_POINT = rf"(?:{_MON}?(?:19|20)\d{{2}}|\d{{1,2}}[./](?:19|20)\d{{2}})"
_NOW = r"(?:present|now|current|today|дотепер|теперішній час|по теперішній|зараз|сьогодні)"
_RANGE = re.compile(rf"(?P<s>{_POINT})\s*(?:[–—-]|to|по|до)\s*(?P<e>{_POINT}|{_NOW})", re.I)
STRONG_VERBS = (
    "built", "launched", "led", "reduced", "increased", "improved", "designed", "delivered", "created", "automated",
    "migrated", "optimised", "optimized", "scaled", "introduced", "drove", "grew", "cut", "saved", "shipped", "owned",
    "implemented", "developed", "managed", "negotiated", "mentored", "redesigned", "rebuilt", "established", "achieved",
    "maintain", "maintained", "delivered", "wrote", "set",
    "розробив", "розробила", "запустив", "запустила", "очолив", "очолила", "скоротив", "скоротила", "збільшив",
    "збільшила", "покращив", "покращила", "впровадив", "впровадила", "автоматизував", "автоматизувала", "побудував",
    "побудувала", "створив", "створила", "оптимізував", "оптимізувала", "мігрував", "мігрувала",
)
_PASSIVE = re.compile(r"\b(was|were|been|being)\s+\w+ed\b", re.I)


def _issue(cat, sev, msg, fix, auto=False):
    return {"category": cat, "severity": sev, "message": msg, "fix": fix, "auto_fixable": auto}


def _today() -> date:
    return date.fromisoformat(os.environ["JOBS_TODAY"]) if os.environ.get("JOBS_TODAY") else date.today()


# ---------------------------------------------------------------- timeline

def _point(text: str) -> tuple[int, int] | None:
    text = text.strip().lower()
    if re.fullmatch(_NOW, text, re.I):
        t = _today()
        return t.year, t.month
    m = re.search(r"(\d{1,2})[./]((?:19|20)\d{2})", text)
    if m:
        return int(m.group(2)), max(1, min(12, int(m.group(1))))
    y = _YEAR.search(text)
    if not y:
        return None
    mon = next((v for k, v in _MONTHS.items() if re.search(rf"\b{k}", text)), None)
    return int(y.group(0)), mon or 1


def _section_of(line: str, current: str) -> str:
    low = line.strip().lower().strip("#*: ")
    if 0 < len(low) <= 40:
        for key, aliases in SECTION_ALIASES.items():
            if any(low == a or low.startswith(a) for a in aliases):
                return key
    return current


def timeline(text: str) -> dict:
    """Positions with dates (outside Education), total experience (overlaps merged), gaps and short tenures."""
    section, positions = "", []
    for line in text.splitlines():
        section = _section_of(line, section)
        if section == "education":
            continue
        for m in _RANGE.finditer(line):
            s, e = _point(m.group("s")), _point(m.group("e"))
            if s and e and e >= s:
                months = (e[0] - s[0]) * 12 + (e[1] - s[1]) + 1
                if months <= 12 * 45:
                    positions.append({"line": re.sub(r"[*#]", "", line).strip()[:90], "start": s, "end": e, "months": months})
    spans = sorted((p["start"][0] * 12 + p["start"][1], p["end"][0] * 12 + p["end"][1]) for p in positions)
    merged: list[list[int]] = []
    for a, b in spans:
        if merged and a <= merged[-1][1] + 1:
            merged[-1][1] = max(merged[-1][1], b)
        else:
            merged.append([a, b])
    total = sum(b - a + 1 for a, b in merged)
    now = _today().year * 12 + _today().month
    fmt = lambda v: f"{(v - 1) // 12}-{(v - 1) % 12 + 1:02d}"  # noqa: E731
    gaps = [{"months": merged[i + 1][0] - merged[i][1] - 1, "between": f"{fmt(merged[i][1])} … {fmt(merged[i + 1][0])}"}
            for i in range(len(merged) - 1) if merged[i + 1][0] - merged[i][1] - 1 >= 6 and now - merged[i + 1][0] < 120]
    short = [p for p in positions if p["months"] < 10 and (p["end"][0] * 12 + p["end"][1]) < now - 1
             and not re.search(r"intern|стаж|trainee|practice|практик", p["line"], re.I)]
    return {"total_years": round(total / 12, 1), "positions": positions, "gaps": gaps, "short_tenures": short}


# ---------------------------------------------------------------- what ATS / recruiters see

def _sections(text: str) -> set[str]:
    found, current = set(), ""
    for line in text.splitlines():
        new = _section_of(line, current)
        if new != current:
            found.add(new)
            current = new
    return found


def terms(text: str) -> set[str]:
    low = text.lower()
    return {t for t in TECH if re.search(rf"(?<![\w.+#-]){re.escape(t)}(?![\w+#-])", low)}


def ats_name_issue(view: dict) -> list[dict]:
    name = (view.get("name") or "").strip()
    bad = (not name or "@" in name or re.search(r"\d{3,}", name) or len(name.split()) > 5
           or _section_of(name, "") != "" or name.lower() in ("cv", "resume", "резюме", "curriculum vitae", "contacts", "контакти"))
    if bad:
        return [_issue("ats", "major", f"An ATS would read '{name[:40]}' as your name (it takes the first line of the text).",
                       "Put your full name as the very first line, then the title and contacts; avoid sidebars before the name.", True)]
    return []


def ats_view(text: str, tl: dict) -> dict:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    email = re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", text)
    phone = re.search(r"\+?\d[\d\s()-]{8,}\d", text)
    linkedin = re.search(r"linkedin\.com/in/[\w-]+", text, re.I)
    return {"name": lines[0].lstrip("# ").strip() if lines else None, "email": email.group(0) if email else None,
            "phone": phone.group(0).strip() if phone else None, "linkedin": linkedin.group(0) if linkedin else None,
            "sections": sorted(_sections(text)), "positions_with_dates": len(tl["positions"]),
            "skills_found": sorted(terms(text)), "years_of_experience": tl["total_years"]}


def _top_block(text: str) -> str:
    """Roughly what fits in the top third of page 1: header + summary, before the first experience entry."""
    out, chars = [], 0
    for line in text.splitlines():
        if _section_of(line, "") in ("experience", "skills") and out:
            break
        out.append(line)
        chars += len(line)
        if chars > 750:
            break
    return "\n".join(out)


def recruiter_scan(text: str, job_title: str | None) -> tuple[list[dict], dict]:
    top = _top_block(text)
    head = [line.strip() for line in top.splitlines() if line.strip()]
    title_line = head[1] if len(head) > 1 else ""
    sees = {
        "title": re.split(r"\s[·|•]\s", title_line)[0][:80] if title_line else None,
        "years_stated": bool(re.search(r"\d+(?:[.,]\d)?\+?\s*(years?|yrs|роки|років|рік)", top, re.I)),
        "key_skills_in_top": sorted(terms(top))[:10],
        "metric_in_top": bool(_METRIC.search(top)),
    }
    issues = []
    if not sees["title"] or len(sees["title"]) < 4 or "@" in sees["title"]:
        issues.append(_issue("first_impression", "major", "No professional title under the name.",
                             "Put the target role right under the name, e.g. 'Senior Project Manager · Fintech'."))
    elif job_title:
        want = set(re.findall(r"[a-zа-яіїєґ]+", job_title.lower())) - {"senior", "middle", "junior", "lead", "the", "and", "of"}
        have = set(re.findall(r"[a-zа-яіїєґ]+", sees["title"].lower()))
        if want and not (want & have):
            issues.append(_issue("first_impression", "major", f"Title '{sees['title']}' does not match the vacancy '{job_title}'.",
                                 "Recruiters and ATS search by job title: if it is truthful, use the vacancy's wording in the header."))
    if not sees["years_stated"]:
        issues.append(_issue("first_impression", "minor", "Years of experience are not stated at the top.",
                             "Open the summary with them: '6 years in project management…'."))
    if len(sees["key_skills_in_top"]) < 3:
        issues.append(_issue("first_impression", "minor", "Fewer than 3 key skills visible in the top third.",
                             "Mention the 3-5 skills that matter most for the target role in the summary."))
    if not sees["metric_in_top"]:
        issues.append(_issue("first_impression", "minor", "No standout number in the top third.",
                             "Lead with the strongest result (e.g. 'cut release time by 40%') — the eye catches numbers first."))
    summary_words = len(re.findall(r"\w+", "\n".join(head[3:])))
    if summary_words > 90:
        issues.append(_issue("first_impression", "minor", f"Summary is long ({summary_words} words).", "Keep it to 2-4 lines."))
    return issues, {"top_block": top[:600], "recruiter_sees": sees}


# ---------------------------------------------------------------- file / ATS readability

def file_checks(path: Path) -> tuple[list[dict], dict]:
    issues, info = [], {"file": path.name, "format": path.suffix.lower().lstrip(".")}
    name = path.name
    if re.search(r"\.(pdf|docx?)\.(pdf|docx?)$", name, re.I):
        issues.append(_issue("ats", "minor", f"Double extension in the file name: '{name}'.",
                             "Use a clear name like 'Name_Surname_Role_CV.pdf' — recruiters see it.", True))
    elif not re.search(r"cv|resume|резюме", name, re.I):
        issues.append(_issue("ats", "minor", f"File name '{name}' does not say it is a CV.", "Use 'Name_Surname_Role_CV.pdf'.", True))
    if path.stat().st_size > 5 * 1024 * 1024:
        issues.append(_issue("ats", "major", "File is larger than 5 MB; many forms reject it.", "Re-export without heavy images.", True))
    fmt = info["format"]
    if fmt == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        pages = len(reader.pages)
        text = "\n".join((p.extract_text() or "") for p in reader.pages)
        info.update(pages=pages, chars=len(text))
        if len(text.strip()) < 150 * pages:
            issues.append(_issue("ats", "critical", "The PDF has (almost) no text layer — it is a scan or an image.",
                                 "ATS and recruiters' search cannot read it. Re-create the CV as a text PDF (render_cv / create_cv).", True))
        try:
            layout = "\n".join((p.extract_text(extraction_mode="layout") or "") for p in reader.pages)
            lines = [line for line in layout.splitlines() if len(line.strip()) > 20]
            split = [line for line in lines if re.search(r"\S {8,}\S", line.strip())]
            if lines and len(split) / len(lines) > 0.3:
                info["layout"] = "multi-column"
                issues.append(_issue("ats", "major", "Two-column layout: ATS often read columns in the wrong order and mix sections.",
                                     "Use a single-column layout (all render_cv themes are single-column).", True))
        except Exception:  # noqa: BLE001 — layout mode is best-effort
            pass
        images = sum(len(p.images) for p in reader.pages[:2]) if pages else 0
        info["images"] = images
        if images:
            issues.append(_issue("ats", "minor", f"{images} image(s) on the first pages (photo, icons or graphics).",
                                 "Fine for Ukrainian companies; for international ones and ATS prefer no photo and text instead of icons.", False))
        if pages > 2:
            issues.append(_issue("readability", "major", f"{pages} pages.", "Keep it to 1-2 pages: cut old / irrelevant experience.", True))
    elif fmt == "docx":
        from docx import Document
        doc = Document(str(path))
        info["tables"] = len(doc.tables)
        if doc.tables:
            issues.append(_issue("ats", "major", "Layout built with tables: ATS may scramble the order.", "Use plain paragraphs and headings.", True))
        header_text = " ".join(p.text for s in doc.sections for p in s.header.paragraphs)
        if re.search(r"@|\+\d", header_text):
            issues.append(_issue("ats", "major", "Contacts are in the page header; many ATS skip headers.",
                                 "Move email / phone into the body of the first page.", True))
    elif fmt in ("md", "txt"):
        issues.append(_issue("ats", "minor", "Plain text / markdown file: it has no design.", "Render a designed PDF (render_cv).", True))
    return issues, info


# ---------------------------------------------------------------- design, measured on the PDF

def design_metrics(pdf: Path) -> tuple[list[dict], dict]:
    """Font sizes / families, margins, text density and fill of the last page, measured from the rendered PDF."""
    issues, m = [], {}
    try:
        import pypdfium2 as pdfium
        import pypdfium2.raw as raw
        doc = pdfium.PdfDocument(str(pdf))
        sizes, ink, fills, margins = [], [], [], []
        for i in range(min(len(doc), 3)):
            page = doc[i]
            tp = page.get_textpage()
            n = tp.count_chars()
            for k in range(0, n, max(1, n // 400)):
                if chr(raw.FPDFText_GetUnicode(tp.raw, k)).isspace():
                    continue
                rect = raw.FS_RECTF()
                if raw.FPDFText_GetLooseCharBox(tp.raw, k, rect):
                    h_pt = rect.top - rect.bottom          # ≈ font size in pt, whatever the producer's matrices
                    if 3 < h_pt < 72:
                        sizes.append(round(h_pt, 1))
            img = page.render(scale=0.5, grayscale=True).to_pil().convert("L")
            w, h = img.size
            px = img.load()
            dark_rows, xmin, xmax, dark = [], w, 0, 0
            for y in range(0, h, 2):
                row_dark = False
                for x in range(0, w, 2):
                    if px[x, y] < 200:
                        dark += 1
                        row_dark = True
                        xmin, xmax = min(xmin, x), max(xmax, x)
                if row_dark:
                    dark_rows.append(y)
            ink.append(dark / ((w // 2) * (h // 2)))
            if dark_rows:
                fills.append(max(dark_rows) / h)
                margins.append((xmin / w * 210, (w - xmax) / w * 210, min(dark_rows) / h * 297))
        pages = len(doc)
        body = sorted(sizes)[len(sizes) // 2] if sizes else None
        m = {"pages": pages, "body_font_pt": body, "min_font_pt": min(sizes) if sizes else None,
             "font_sizes": sorted(set(round(s) for s in sizes)), "ink_density": round(sum(ink) / len(ink), 3) if ink else None,
             "last_page_fill": round(fills[-1], 2) if fills else None,
             "margins_mm": [round(x) for x in margins[0]] if margins else None}
        try:
            from pypdf import PdfReader
            fonts = set()
            for p in PdfReader(str(pdf)).pages[:2]:
                for f in ((p.get("/Resources") or {}).get("/Font") or {}).values():
                    base = str(f.get_object().get("/BaseFont", ""))
                    fonts.add(re.sub(r"^/?[A-Z]{6}\+", "", base).lstrip("/").split("-")[0].split(",")[0])
            m["font_families"] = sorted(x for x in fonts if x)
        except Exception:  # noqa: BLE001
            pass
        if body and body < 9:
            issues.append(_issue("design", "major", f"Body text is {body} pt — hard to read.", "Use 10-11 pt for body text.", True))
        elif m["min_font_pt"] and m["min_font_pt"] < 7.5:
            issues.append(_issue("design", "minor", f"Some text is {m['min_font_pt']} pt.", "Nothing below 8 pt.", True))
        if len(m.get("font_families", [])) > 3:
            issues.append(_issue("design", "minor", f"{len(m['font_families'])} different fonts: {', '.join(m['font_families'][:5])}.",
                                 "Use one font family (two at most).", True))
        if m["margins_mm"] and min(m["margins_mm"][:2]) < 9:
            issues.append(_issue("design", "major", f"Side margins are ~{min(m['margins_mm'][:2])} mm — the page looks cramped.",
                                 "Keep 12-20 mm margins; white space makes a CV readable.", True))
        if m["ink_density"] and m["ink_density"] > 0.2:
            issues.append(_issue("design", "major", "Very dense page: a wall of text.", "Shorter bullets, more spacing between sections.", True))
        elif m["ink_density"] is not None and m["ink_density"] < 0.03:
            issues.append(_issue("design", "minor", "Page is mostly empty.", "Add achievements or use a compact one-page layout.", True))
        if pages > 1 and m["last_page_fill"] is not None and m["last_page_fill"] < 0.25:
            issues.append(_issue("design", "minor", f"The last page is only {int(m['last_page_fill'] * 100)}% filled.",
                                 "Fit it to one page (render_cv max_pages=1) or balance the content.", True))
    except Exception as e:  # noqa: BLE001 — design metrics are optional
        m["error"] = f"{type(e).__name__}: {e}"
    return issues, m


# ---------------------------------------------------------------- content

def _bullets(text: str) -> list[str]:
    found = [_BULLET.sub("", line).strip() for line in text.splitlines() if _BULLET.match(line) and len(line.strip()) > 12]
    if found:
        return found
    # PDFs often draw bullets as graphics: take sentence-like lines inside the Experience section
    section, out = "", []
    for line in text.splitlines():
        section = _section_of(line, section)
        stripped = line.strip()
        if section == "experience" and len(stripped.split()) >= 6 and not _RANGE.search(stripped) \
                and not any(re.search(rf"\b{re.escape(w)}\b", stripped.lower()) for w in ROLE_WORDS[:12] if " — " in stripped):
            out.append(stripped)
    return out if len(out) >= 3 else []


def _job_lines(text: str) -> list[tuple[str, bool]]:
    lines = [line.strip() for line in text.splitlines()]
    out = []
    for i, line in enumerate(lines):
        low = line.lower()
        if not (8 < len(line) < 160) or _BULLET.match(line) or line.startswith("##"):
            continue
        looks_role = any(re.search(rf"\b{re.escape(w)}\b", low) for w in ROLE_WORDS)
        has_company_sep = bool(re.search(r"\s[—–-]\s|\s@\s|\bat\b|,\s", line))
        if looks_role and has_company_sep and (line.startswith("**") or line.startswith("###") or _YEAR.search(line)
                                               or (i + 1 < len(lines) and _YEAR.search(lines[i + 1]))):
            window = line + " " + (lines[i + 1] if i + 1 < len(lines) else "")
            out.append((line.strip("*# "), bool(_RANGE.search(window) or _YEAR.search(window))))
    return out


def content_checks(text: str, tl: dict) -> tuple[list[dict], dict]:
    issues, low = [], text.lower()
    words = len(re.findall(r"\w+", text))
    info = {"words": words}
    if not re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", text):
        issues.append(_issue("completeness", "critical", "No email.", "Add an email at the top."))
    first = next((line.lstrip("# ").strip() for line in text.splitlines() if line.strip()), "")
    latin = len(re.findall(r"[A-Za-z]", text)) / max(1, len(re.findall(r"[A-Za-zА-Яа-яІіЇїЄєҐґ]", text)))
    if re.search(r"[А-Яа-яІіЇїЄєҐґ]", first) and latin > 0.7:
        from jobs_helper.cv.naming import translit
        issues.append(_issue("ats", "minor", f"The CV is in English but the name is in Cyrillic ('{first}').",
                             f"Write it in Latin as in the passport (e.g. '{translit(first)}') — recruiters search by it."))
    if not re.search(r"\+?\d[\d\s()-]{8,}\d", text):
        issues.append(_issue("completeness", "major", "No phone number.", "Add a phone in +380… format."))
    if "linkedin.com" not in low:
        issues.append(_issue("completeness", "minor", "No LinkedIn link.", "Add linkedin.com/in/… next to the contacts."))
    secs = _sections(text)
    info["sections"] = sorted(secs)
    for key, sev in (("experience", "critical"), ("skills", "major"), ("summary", "minor"), ("education", "minor"),
                     ("languages", "minor")):
        if key not in secs:
            issues.append(_issue("completeness", sev, f"No '{key}' section.", f"Add a clearly titled '{key.capitalize()}' section."))
    undated = [j for j, dated in _job_lines(text) if not dated]
    if undated:
        issues.append(_issue("completeness", "major", f"No dates for: {'; '.join(u[:60] for u in undated[:4])}.",
                             "Add month/year – month/year for every position; recruiters skip CVs with unclear timelines."))
    for g in tl["gaps"][:2]:
        issues.append(_issue("completeness", "minor", f"Gap of {g['months']} months ({g['between']}).",
                             "Explain it in one line if relevant (studies, freelance, relocation, military service)."))
    if len(tl["short_tenures"]) >= 2:
        issues.append(_issue("completeness", "minor", f"{len(tl['short_tenures'])} positions shorter than 10 months.",
                             "Group short contracts under one 'Freelance / Contract' entry or explain them."))
    bullets = _bullets(text)
    info["bullets"] = len(bullets)
    if bullets:
        quantified = [b for b in bullets if _NUMBER.search(b)]
        share = len(quantified) / len(bullets)
        info["quantified_share"] = round(share, 2)
        if share < 0.3:
            issues.append(_issue("impact", "major", f"Only {len(quantified)} of {len(bullets)} bullets have numbers.",
                                 "Add measurable results where known: %, time saved, users, revenue, team size. Never invent numbers."))
        elif share < 0.5:
            issues.append(_issue("impact", "minor", f"{len(quantified)} of {len(bullets)} bullets have numbers.",
                                 "Top CVs quantify about half of the bullets."))
        weak = [b for b in bullets if b.lower().startswith(WEAK_STARTS)]
        if weak:
            issues.append(_issue("impact", "major" if len(weak) > 2 else "minor",
                                 f"{len(weak)} bullet(s) describe duties, not results: “{weak[0][:70]}…”.",
                                 "Start with an action verb and the outcome: 'Built…', 'Reduced…', 'Launched…'."))
        exp_bullets = [b for b in bullets if not re.match(r"^\*\*[^*]+:\*\*", b)]   # skip 'Languages: …' skill lines
        strong = [b for b in exp_bullets if b.split()[0].lower().strip(",.") in STRONG_VERBS]
        info["strong_verb_share"] = round(len(strong) / len(exp_bullets), 2) if exp_bullets else None
        if exp_bullets and len(strong) / len(exp_bullets) < 0.4 and not weak:
            issues.append(_issue("impact", "minor", "Few bullets start with a strong action verb.",
                                 "Lead with verbs like Built, Launched, Reduced, Led, Automated."))
        firsts = [b.split()[0].lower() for b in exp_bullets]
        repeated = sorted({w for w in firsts if firsts.count(w) >= 3 and w.isalpha()})
        if repeated:
            issues.append(_issue("readability", "minor", f"The same opening verb repeats: {', '.join(repeated)}.", "Vary the verbs."))
        long_b = [b for b in bullets if len(b.split()) > 32]
        if long_b:
            issues.append(_issue("readability", "minor", f"{len(long_b)} bullet(s) longer than 32 words.", "Split or shorten to 1-2 lines."))
    elif "experience" in secs:
        issues.append(_issue("impact", "major", "Experience is written as paragraphs, without bullet points.",
                             "Use 2-5 bullets per position, each with an action and a result."))
    if len(_PASSIVE.findall(text)) >= 2:
        issues.append(_issue("impact", "minor", "Passive voice ('was implemented…').", "Say who did it: 'Implemented…'."))
    cliches = sorted({c for c in CLICHES if c in low})
    if cliches:
        issues.append(_issue("impact", "minor", f"Clichés: {', '.join(cliches[:5])}.", "Replace with evidence (what you did and the result)."))
    if re.search(r"(^|\n)\s*[-•*]?\s*(i|my|я|мій|моя)\b", low):
        issues.append(_issue("readability", "minor", "First-person sentences in bullets.", "Drop 'I/my': start with the verb."))
    if words < 150:
        issues.append(_issue("readability", "major", f"Very short CV ({words} words).", "Describe achievements for each role."))
    elif words > 950:
        issues.append(_issue("readability", "major", f"Long CV ({words} words).", "Aim for 350-700 words.", True))
    return issues, info


# ---------------------------------------------------------------- vacancy filters

def match_checks(cv_text: str, job_text: str, tl: dict, job: dict | None = None) -> tuple[list[dict], dict, list[dict]]:
    job = job or {}
    title = job.get("title") or job_text.splitlines()[0]
    title_terms = terms(title)
    wanted, have = terms(job_text), terms(cv_text)
    weight = {t: (2 if t in title_terms else 1) for t in wanted}
    covered = sorted(wanted & have)
    missing = sorted(wanted - have, key=lambda t: (-weight[t], t))
    coverage = (sum(weight[t] for t in covered) / sum(weight.values())) if wanted else None
    info = {"vacancy_keywords": sorted(wanted), "covered": covered, "missing": missing,
            "coverage": round(coverage, 2) if coverage is not None else None}
    issues, knockouts = [], []
    req = job.get("experience_years")
    if req is None:
        m = re.search(r"(\d{1,2})\+?\s*(?:years?|yrs|роки|років)", job_text, re.I)
        req = float(m.group(1)) if m else None
    info["required_years"], info["candidate_years"] = req, tl["total_years"]
    if req and not tl["total_years"]:
        knockouts.append({"filter": "experience", "required": req, "cv_shows": "no dates"})
        issues.append(_issue("match", "critical", f"The vacancy asks for {req:g}+ years, but the CV has no job dates — an ATS "
                             "cannot compute any experience and filters the CV out.", "Add month/year – month/year to every position."))
    elif req and tl["total_years"] + 0.5 < req:
        knockouts.append({"filter": "experience", "required": req, "cv_shows": tl["total_years"]})
        issues.append(_issue("match", "critical", f"The vacancy asks for {req:g}+ years, the CV dates show {tl['total_years']:g}.",
                             "ATS filters on years: apply only if other experience is missing from the CV (add it with dates), "
                             "otherwise target a level lower."))
    from jobs_helper.extract import guess_english
    want_en, have_en = guess_english(job_text), guess_english(cv_text)
    order = ["A1", "A2", "B1", "B2", "C1", "C2"]
    if want_en in order and have_en in order and order.index(have_en) < order.index(want_en):
        knockouts.append({"filter": "english", "required": want_en, "cv_shows": have_en})
        issues.append(_issue("match", "major", f"English {want_en} is required, the CV says {have_en}.",
                             "If your level is higher, state it precisely (e.g. 'Upper-Intermediate, B2')."))
    elif want_en and not have_en:
        issues.append(_issue("match", "minor", f"The vacancy requires English ({want_en}) and the CV does not state a level.",
                             "Add a Languages line with the CEFR level."))
    if wanted and coverage is not None and coverage < 0.6:
        knockouts.append({"filter": "keywords", "coverage": round(coverage, 2), "missing": missing[:8]})
        issues.append(_issue("match", "major", f"Key technologies covered: {int(coverage * 100)}% "
                             f"(CV mentions {len(covered)} of {len(wanted)}).",
                             f"If the candidate really has them, make them visible (skills + a bullet): {', '.join(missing[:8])}. "
                             "If not — do NOT add them; mention learning plans in the cover letter."))
    elif missing:
        issues.append(_issue("match", "minor", f"Not mentioned: {', '.join(missing[:8])}.", "Add only those the candidate actually used."))
    return issues, info, knockouts


def unverified_claims(new_text: str, sources: str) -> dict:
    """Numbers and technologies in a rewritten CV that appear nowhere in the candidate's own documents / notes."""
    src_numbers = set(re.findall(r"\d+(?:[.,]\d+)?", sources))
    nums = sorted({n for n in re.findall(r"\d+(?:[.,]\d+)?", new_text) if n not in src_numbers and not _YEAR.fullmatch(n)})
    return {"numbers": nums, "technologies": sorted(terms(new_text) - terms(sources))}


# ---------------------------------------------------------------- report

def review(text: str, file: Path | None = None, job_text: str | None = None, job: dict | None = None,
           design_pdf: Path | None = None) -> dict:
    issues, info, knockouts = [], {}, []
    tl = timeline(text)
    if file:
        fi, finfo = file_checks(file)
        issues += fi
        info["file"] = finfo
        if any(i["severity"] == "critical" and "text layer" in i["message"] for i in fi):
            score = max(0, 100 - sum(PENALTY[i["severity"]] for i in fi))
            return {"overall": min(35, score), "level": "Bronze", "grade": "weak", "scores": {"ats": score}, "issues": fi,
                    "action_plan": [{"fix": i["fix"], "gain": 30, "severity": i["severity"], "category": i["category"]} for i in fi[:5]],
                    "top_fixes": [i["fix"] for i in fi[:5]], "details": info, "design_fixable": True, "knockouts": [],
                    "note": "Content was not checked: the file has no text. Ask for a text version or rebuild it (create_cv)."}
        if design_pdf is None and file.suffix.lower() == ".pdf":
            design_pdf = file
    ci, cinfo = content_checks(text, tl)
    issues += ci
    info["content"] = cinfo
    ri, rinfo = recruiter_scan(text, (job or {}).get("title"))
    issues += ri
    issues += ats_name_issue(ats_view(text, tl))
    cats = ["ats", "completeness", "first_impression", "impact", "readability"]
    if design_pdf:
        di, dinfo = design_metrics(design_pdf)
        issues += di
        info["design"] = dinfo
        cats.append("design")
    if job_text:
        mi, minfo, knockouts = match_checks(text, job_text, tl, job)
        issues += mi
        info["match"] = minfo
        cats.append("match")
    scores = {c: max(0, 100 - sum(PENALTY[i["severity"]] for i in issues if i["category"] == c)) for c in cats}
    total_w = sum(WEIGHTS[c] for c in cats)
    overall = round(sum(scores[c] * WEIGHTS[c] for c in cats) / total_w)
    n_major = sum(i["severity"] == "major" for i in issues)
    if any(i["severity"] == "critical" for i in issues):
        overall = min(overall, 55)
    elif n_major >= 4:
        overall = min(overall, 65)
    elif n_major >= 3:
        overall = min(overall, 72)
    elif n_major >= 2:
        overall = min(overall, 79)
    for i in issues:   # points this fix is worth for the overall score (before caps)
        i["gain"] = max(1, round(PENALTY[i["severity"]] * WEIGHTS[i["category"]] / total_w))
    order = {"critical": 0, "major": 1, "minor": 2}
    issues.sort(key=lambda i: (order[i["severity"]], -i["gain"]))
    level, label = next((name, text_) for th, name, text_ in LEVELS if overall >= th)
    nxt = next(((th, name) for th, name, _ in reversed(LEVELS) if th > overall), None)
    plan = [{"fix": i["fix"], "gain": i["gain"], "severity": i["severity"], "category": i["category"]} for i in issues[:8]]
    grade = "strong" if overall >= 88 else "good" if overall >= 75 else "needs work" if overall >= 55 else "weak"
    return {
        "overall": overall, "level": level, "level_meaning": label, "grade": grade,
        "next_level": {"name": nxt[1], "points_needed": nxt[0] - overall} if nxt else None,
        "scores": scores, "knockouts": knockouts,
        "recruiter_view": rinfo, "ats_view": ats_view(text, tl),
        "timeline": {"total_years": tl["total_years"], "gaps": tl["gaps"],
                     "short_tenures": [p["line"] for p in tl["short_tenures"]]},
        "action_plan": plan, "issues": issues, "top_fixes": [i["fix"] for i in issues[:5]], "details": info,
        "design_fixable": any(i["auto_fixable"] and i["category"] in ("ats", "readability", "design") for i in issues),
    }
