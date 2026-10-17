"""Per-site submit policy chosen by the human (settings.yaml) and the candidate profile (candidate.yaml).

The agent never chooses how an application is sent: the MCP server reads this file itself."""

from __future__ import annotations

import json
import os
import re
from pathlib import Path
from typing import Literal

import yaml

from jobs_helper import config

Policy = Literal["manual", "confirm", "auto"]
POLICIES: tuple[Policy, ...] = ("manual", "confirm", "auto")
SITES = ("dou", "djinni", "happymonday", "ats", "linkedin")
DEFAULTS: dict[str, Policy] = {"dou": "manual", "djinni": "confirm", "happymonday": "confirm", "ats": "confirm",
                               "linkedin": "confirm"}

SITE_NOTES = {
    "dou": "Правила DOU забороняють «інструменти для автоматичних відгуків» і масові відгуки "
           "(dou.ua/applications-rules); за порушення можуть заблокувати профіль.",
    "djinni": "Окремої заборони автоматизації для кандидатів не знайдено.",
    "happymonday": "Окремої заборони автоматизації для кандидатів не знайдено.",
    "ats": "Сайти компаній / ATS: можлива капча; правила DOU тут не діють.",
    "linkedin": "Публікація через офіційне API (Share on LinkedIn).",
}


def settings_path() -> Path:
    return Path(os.environ.get("JOBS_SETTINGS", config.HOME / "settings.yaml"))


def candidate_path() -> Path:
    return Path(os.environ.get("JOBS_CANDIDATE", config.HOME / "candidate.yaml"))


def _load_yaml(path: Path) -> dict:
    try:
        return yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    except FileNotFoundError:
        return {}


def _truthy(var: str) -> bool:
    return os.environ.get(var, "").strip().lower() in ("1", "true", "yes", "так")


def auto_toggle(site: str) -> bool:
    """'Send automatically' switch of the Claude Desktop extension form."""
    return _truthy(f"JOBS_AUTO_{site.upper()}")


def load_policies() -> dict[str, Policy]:
    """Order: extension toggle 'auto' > settings.yaml (chosen in chat or CLI) > JOBS_POLICY_<SITE> > defaults."""
    raw = _load_yaml(settings_path()).get("submit_policy", {}) or {}
    out = {}
    for site in SITES:
        env = os.environ.get(f"JOBS_POLICY_{site.upper()}", "").strip().lower()
        if auto_toggle(site):
            value = "auto"
        elif raw.get(site) in POLICIES:
            value = raw[site]
        else:
            value = env
        out[site] = value if value in POLICIES else DEFAULTS[site]
    return out


def explicit_policies() -> dict[str, str]:
    """Only the sites the human has actually chosen for."""
    raw = _load_yaml(settings_path()).get("submit_policy", {}) or {}
    return {k: v for k, v in raw.items() if v in POLICIES}


def set_policy(site: str, policy: Policy) -> None:
    if site not in SITES or policy not in POLICIES:
        raise ValueError(f"site must be one of {SITES}, policy one of {POLICIES}")
    path = settings_path()
    data = _load_yaml(path)
    data.setdefault("submit_policy", {})[site] = policy
    path.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")


def dry_run() -> bool:
    return os.environ.get("JOBS_DRY_RUN", "").lower() in ("1", "true", "yes")


# environment variable -> (section, key); used by the Claude Desktop extension form
ENV_PROFILE = {
    "JOBS_CANDIDATE_FIRST_NAME": ("personal", "first_name"), "JOBS_CANDIDATE_LAST_NAME": ("personal", "last_name"),
    "JOBS_CANDIDATE_EMAIL": ("personal", "email"), "JOBS_CANDIDATE_PHONE": ("personal", "phone"),
    "JOBS_CANDIDATE_CITY": ("personal", "city"), "JOBS_CANDIDATE_COUNTRY": ("personal", "country"),
    "JOBS_CANDIDATE_LINKEDIN": ("links", "linkedin"), "JOBS_CANDIDATE_GITHUB": ("links", "github"),
    "JOBS_CANDIDATE_PORTFOLIO": ("links", "portfolio"), "JOBS_CANDIDATE_TELEGRAM": ("links", "telegram"),
    "JOBS_CANDIDATE_SALARY_USD": ("job_preferences", "desired_salary_usd"),
    "JOBS_CANDIDATE_NOTICE": ("job_preferences", "notice_period"),
    "JOBS_CANDIDATE_ENGLISH": ("languages", "english"),
    "JOBS_CANDIDATE_CV": (None, "cv_file"),
    "JOBS_CANDIDATE_CONSENT": ("answers", "consent_to_processing"),
}


def load_candidate() -> dict | None:
    """candidate.yaml overlaid with JOBS_CANDIDATE_* variables; empty values are ignored."""
    data = _load_yaml(candidate_path())
    for var, (section, key) in ENV_PROFILE.items():
        raw = os.environ.get(var, "").strip()
        if not raw or raw.startswith("${"):          # unset extension fields may arrive as literal placeholders
            continue
        value: object = raw
        if key == "desired_salary_usd":
            try:
                value = int(float(raw))
            except ValueError:
                continue
        if key == "consent_to_processing":
            value = raw.lower() in ("1", "true", "yes", "так")
        (data.setdefault(section, {}) if section else data)[key] = value
    links = data.get("links") or {}
    for k in ("linkedin", "github", "portfolio"):
        v = str(links.get(k) or "").strip()
        if v and not v.startswith(("http://", "https://")) and "." in v:
            links[k] = "https://" + v.lstrip("/")
    saved = {n["question"]: n["text"] for n in list_notes() if n.get("kind") == "answer" and n.get("question")}
    if saved:
        data["saved_answers"] = saved
    return data or None


CV_NAMES = ("cv.md", "cv.txt", "cv.pdf", "cv.docx")
DOC_SUFFIXES = (".pdf", ".docx", ".md", ".txt")
_CV_NAME = re.compile(r"\bcv\b|resume|résumé|резюме|curriculum", re.I)


def candidate_files(profile: dict | None = None) -> list[Path]:
    """Every document the candidate gave us: JOBS_CANDIDATE_FILES (extension form, several files), cv_file, ~/.jobs_helper/cv.*"""
    profile = profile if profile is not None else (load_candidate() or {})
    found: list[Path] = []
    raw = [x for x in os.environ.get("JOBS_CANDIDATE_FILES", "").split(os.pathsep) if x.strip() and not x.startswith("${")]
    if profile.get("cv_file"):
        p = Path(str(profile["cv_file"])).expanduser()
        raw.insert(0, str(p if p.is_absolute() else candidate_path().parent / p))
    raw += [str(config.HOME / n) for n in CV_NAMES]
    for item in raw:
        path = Path(item).expanduser()
        if path.exists() and path.suffix.lower() in DOC_SUFFIXES and path.resolve() not in [f.resolve() for f in found]:
            found.append(path)
    return found


def cv_file(profile: dict | None = None) -> Path | None:
    """The main CV: an explicit cv_file, else the first file named like a CV, else the first document."""
    profile = profile if profile is not None else (load_candidate() or {})
    files = candidate_files(profile)
    if profile.get("cv_file") and files:
        return files[0]
    return next((f for f in files if _CV_NAME.search(f.stem.replace("_", " "))), files[0] if files else None)


# ------------------------------------------------------------ candidate notes (remembered from the chat)

def notes_path() -> Path:
    return config.HOME / "candidate_notes.json"


def list_notes() -> list[dict]:
    try:
        return json.loads(notes_path().read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []


def add_note(text: str, kind: str = "fact", question: str | None = None) -> dict:
    notes = list_notes()
    note = {"id": f"N{max([int(n['id'][1:]) for n in notes] or [0]) + 1}", "kind": kind, "text": text.strip(),
            "question": (question or "").strip() or None, "created": __import__("datetime").date.today().isoformat()}
    notes.append(note)
    notes_path().parent.mkdir(parents=True, exist_ok=True)
    notes_path().write_text(json.dumps(notes, ensure_ascii=False, indent=1), encoding="utf-8")
    return note


def remove_note(note_id: str) -> bool:
    notes = list_notes()
    kept = [n for n in notes if n["id"] != note_id]
    if len(kept) == len(notes):
        return False
    notes_path().write_text(json.dumps(kept, ensure_ascii=False, indent=1), encoding="utf-8")
    return True


def read_cv_text(path: Path) -> str:
    """Plain text of a CV in .md / .txt / .pdf / .docx."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages).strip()
    if suffix == ".docx":
        from docx import Document
        return "\n".join(p.text for p in Document(str(path)).paragraphs).strip()
    return path.read_text(errors="replace", encoding="utf-8").strip()
