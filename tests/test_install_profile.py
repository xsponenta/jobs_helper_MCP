"""Profile from the Claude Desktop extension form (env), CV reading and the get_candidate tool."""

import json

from docx import Document

from jobs_helper import policy as pol


def test_env_overlays_yaml_and_ignores_placeholders(tmp_path, monkeypatch):
    cand = tmp_path / "candidate.yaml"
    cand.write_text("personal: {first_name: Old, city: Kyiv}\n")
    monkeypatch.setenv("JOBS_CANDIDATE", str(cand))
    monkeypatch.setenv("JOBS_CANDIDATE_FIRST_NAME", "Olena")
    monkeypatch.setenv("JOBS_CANDIDATE_PHONE", "${user_config.phone}")   # unset field of the extension form
    monkeypatch.setenv("JOBS_CANDIDATE_SALARY_USD", "3500.0")
    monkeypatch.setenv("JOBS_CANDIDATE_CONSENT", "false")
    monkeypatch.setenv("JOBS_POLICY_DOU", "confirm")
    monkeypatch.setenv("JOBS_POLICY_ATS", "")                            # empty -> default
    prof = pol.load_candidate()
    assert prof["personal"] == {"first_name": "Olena", "city": "Kyiv"}
    assert prof["job_preferences"]["desired_salary_usd"] == 3500
    assert prof["answers"]["consent_to_processing"] is False
    assert pol.load_policies()["dou"] == "confirm" and pol.load_policies()["ats"] == "confirm"


def test_cv_formats(tmp_path, monkeypatch):
    monkeypatch.setenv("JOBS_CANDIDATE", str(tmp_path / "none.yaml"))
    doc = Document()
    doc.add_paragraph("Olena Kovalenko — Python Backend Developer")
    doc.save(tmp_path / "cv.docx")
    monkeypatch.setenv("JOBS_CANDIDATE_CV", str(tmp_path / "cv.docx"))
    path = pol.cv_file()
    assert path.name == "cv.docx" and "Python Backend" in pol.read_cv_text(path)


def test_get_candidate_tool(tmp_path, monkeypatch):
    (tmp_path / "cv.md").write_text("# Olena\nPython, Django")
    monkeypatch.setenv("JOBS_CANDIDATE", str(tmp_path / "none.yaml"))
    monkeypatch.setenv("JOBS_CANDIDATE_FIRST_NAME", "Olena")
    monkeypatch.setenv("JOBS_CANDIDATE_CV", str(tmp_path / "cv.md"))
    from jobs_helper.mcp_all import get_candidate
    out = get_candidate()
    assert out["profile"]["personal"]["first_name"] == "Olena" and "Django" in out["cv_text"]
    assert "sensitive" not in json.dumps(out["profile"])


def test_several_documents_main_cv_by_name(tmp_path, monkeypatch):
    import os
    monkeypatch.setenv("JOBS_CANDIDATE", str(tmp_path / "none.yaml"))
    (tmp_path / "portfolio.md").write_text("Portfolio: booking app")
    (tmp_path / "Vus_Stanislav_CV_PM.pdf.md").write_text("PM with 5 years")
    monkeypatch.setenv("JOBS_CANDIDATE_FILES", os.pathsep.join([str(tmp_path / "portfolio.md"), str(tmp_path / "Vus_Stanislav_CV_PM.pdf.md"), "${user_config.documents}"]))
    files = pol.candidate_files({})
    assert [f.name for f in files] == ["portfolio.md", "Vus_Stanislav_CV_PM.pdf.md"]
    assert pol.cv_file({}).name == "Vus_Stanislav_CV_PM.pdf.md"


def test_remembered_facts_and_answers(tmp_path, monkeypatch):
    monkeypatch.setattr("jobs_helper.config.HOME", tmp_path)
    monkeypatch.setenv("JOBS_CANDIDATE", str(tmp_path / "none.yaml"))
    monkeypatch.setenv("JOBS_CANDIDATE_FIRST_NAME", "Ann")
    from jobs_helper.mcp_all import forget_candidate_info, get_candidate, remember_candidate_info
    a = remember_candidate_info("Led a team of 5 at PwC (2021-2023)")
    b = remember_candidate_info("I want a product company", kind="answer", question="Why are you looking for a new position?")
    out = get_candidate()
    assert out["remembered_facts"][0]["text"].startswith("Led a team")
    assert pol.load_candidate()["saved_answers"] == {"Why are you looking for a new position?": "I want a product company"}
    assert forget_candidate_info(a["saved"]["id"])["ok"] and len(get_candidate()["remembered_facts"]) == 0
    assert remember_candidate_info("x", kind="answer")["error"] == "invalid_argument"


def test_toggle_beats_chat_choice_and_linkedin_url(tmp_path, monkeypatch):
    monkeypatch.setenv("JOBS_SETTINGS", str(tmp_path / "settings.yaml"))
    monkeypatch.setenv("JOBS_CANDIDATE", str(tmp_path / "none.yaml"))
    from jobs_helper.mcp_all import set_submit_policy
    assert set_submit_policy("djinni", "manual")["policy"] == "manual"
    monkeypatch.setenv("JOBS_AUTO_DJINNI", "true")
    r = set_submit_policy("djinni", "confirm")
    assert r["policy"] == "auto" and "warning" in r
    monkeypatch.setenv("JOBS_CANDIDATE_LINKEDIN", "www.linkedin.com/in/stanislav-vus")
    assert pol.load_candidate()["links"]["linkedin"] == "https://www.linkedin.com/in/stanislav-vus"
